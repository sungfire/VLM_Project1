"""Build report/project1_report.docx (English) from results/ and report/stats.json."""
import json
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from src.data import load_hrbench  # noqa: E402
from src.prompts import JOINT_SUFFIX, TYPE_PROMPT  # noqa: E402
from src.run_eval import load_jsonl  # noqa: E402

S = json.load(open("report/stats.json"))
R = S["runs"]
F1, F2 = S["final_success1_4k"], S["final_success2_2k"]
A1, A2 = S["ablation_4k"], S["ablation_2k"]
LET = "ABCD"


def am(p):
    return LET[p.index(max(p))]


# ---- extra numbers computed from logs
jd = load_jsonl("results/joint_dev.jsonl")
eff = {"single": [0, 0], "cross": [0, 0]}  # zoom fixed, zoom broke
for r in jd:
    if len(r["calls"]) >= 2:
        a1 = am(r["calls"][0]["probs"]) == r["answer"]
        a3 = am(r["calls"][1]["probs"]) == r["answer"]
        eff[r["category"]][0] += (not a1) and a3
        eff[r["category"]][1] += a1 and not a3
hi = [r for r in jd if max(r["calls"][0]["probs"]) >= 0.999]
hi_acc = 100 * sum(am(r["calls"][0]["probs"]) == r["answer"] for r in hi) / len(hi)
lo = [r for r in jd if max(r["calls"][0]["probs"]) < 0.999]
lo_acc = 100 * sum(am(r["calls"][0]["probs"]) == r["answer"] for r in lo) / len(lo)
imgs1 = sum(F1["per_call"][k]["count"] for k in ("probe", "single", "relational")) / 800
imgs2 = sum(F2["per_call"][k]["count"] for k in ("probe", "single")) / 800
Q = {i.idx: i.question.lower() for i in load_hrbench()}
QX = {}
for n, k in [("success1_4k", "s1"), ("success2_2k", "s2")]:
    rows = load_jsonl(f"results/{n}.jsonl")
    QX[f"{k}_map"] = sum(r["correct"] for r in rows if "map" in Q[r["idx"]])
    QX[f"{k}_cnt"] = sum(r["correct"] for r in rows if Q[r["idx"]].startswith("how many"))

# ---------------------------------------------------------------- document setup
doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin = sec.bottom_margin = Cm(2.0)


def set_font(style, size, bold=None):
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rf.set(qn(a), "Calibri")


set_font(doc.styles["Normal"], 11)
doc.styles["Normal"].paragraph_format.space_after = Pt(3)
doc.styles["Normal"].paragraph_format.line_spacing = 1.0
for name, size in [("Title", 16), ("Heading 1", 13), ("Heading 2", 11.5)]:
    set_font(doc.styles[name], size, True)
    doc.styles[name].font.color.rgb = None
    doc.styles[name].paragraph_format.space_before = Pt(8 if name != "Title" else 0)
    doc.styles[name].paragraph_format.space_after = Pt(3)


def para(text, bold_prefix=None, size=None, align=None, italic=False):
    p = doc.add_paragraph()
    if bold_prefix:
        r = p.add_run(bold_prefix)
        r.bold = True
        if size:
            r.font.size = Pt(size)
    r = p.add_run(text)
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    p.paragraph_format.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    return p


def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(1)
    return p


def caption(text):
    p = para(text, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)
    p.paragraph_format.space_after = Pt(6)


def table(header, rows, widths, size=8.5, bold_rows=()):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, gc in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(Cm(widths[i]).twips)))
    for i, h in enumerate(header):
        r = t.rows[0].cells[i].paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(size)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(size)
            r.bold = ri in bold_rows
    for row in t.rows:
        for i, c in enumerate(row.cells):
            c.width = Cm(widths[i])
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            if i > 0:
                c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    return t


def fig(path, width_cm, cap):
    doc.add_picture(path, width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.keep_with_next = True
    caption(cap)


def f1(x):
    return f"{x:.1f}"


def run_row(label, s):
    return [label, f1(s["mean_visual"]), f"{s['acc']:.2f}", f"{s['acc_single']:.2f}", f"{s['acc_cross']:.2f}",
            f"{s['mean_calls']:.2f}"]


def abl_row(label, v):
    return [label, f1(v["tok"]), f"{v['acc']:.2f}", f"{v['single']:.2f}", f"{v['cross']:.2f}"]


# ---------------------------------------------------------------- title
doc.add_paragraph("Efficient High-Resolution VQA under a Visual-Token Budget:\n"
                  "Joint Probe–Zoom Routing for HR-Bench 4K", style="Title")
para("Advanced Topics in Multimodal AI — Midterm Project 1      Student ID: 2026xxxxxx      Name: ________", size=10,
     align=WD_ALIGN_PARAGRAPH.LEFT)

# ---------------------------------------------------------------- 1
doc.add_heading("1. Overview", level=1)
para("We design a training-free inference procedure that answers HR-Bench 4K (200 questions × 4 option orderings = 800 instances) "
     "with a frozen Qwen3-VL-2B-Instruct, while keeping the mean cumulative number of visual tokens per instance below 4,096 "
     "(Success 1) or 2,048 (Success 2). The method has three ingredients: (1) a joint probe that obtains both the answer "
     "distribution and an evidence bounding box from a single low-resolution global view, (2) confidence-based early stopping, and "
     "(3) a text-only question-type classifier that routes uncertain instances either to a zoomed crop or to a higher-resolution "
     "global re-observation.")
table(["Condition", "Mean visual tokens", "Accuracy (%)", "Single", "Cross", "Target (Main / Optional)"],
      [["Success 1", f1(R["success1_4k"]["mean_visual"]) + " ≤ 4,096", f"{R['success1_4k']['acc']:.2f}",
        f"{R['success1_4k']['acc_single']:.2f}", f"{R['success1_4k']['acc_cross']:.2f}", "≥71 ✓ / ≥74 ✓"],
       ["Success 2", f1(R["success2_2k"]["mean_visual"]) + " ≤ 2,048", f"{R['success2_2k']['acc']:.2f}",
        f"{R['success2_2k']['acc_single']:.2f}", f"{R['success2_2k']['acc_cross']:.2f}", "≥68 ✓ / ≥71 ✓"]],
      widths=[2.2, 3.4, 2.2, 1.6, 1.6, 3.8], size=9.5)
caption("Table 1. Final results (all 800 instances, end-to-end runs; a fresh re-run reproduced all 800 predictions and token counts).")

# ---------------------------------------------------------------- 2
doc.add_heading("2. Method", level=1)
doc.add_heading("2.1 Visual-token accounting", level=2)
para("Qwen3-VL uses 16×16 patches with 2×2 spatial merging, so one visual token covers 32×32 pixels. 136 of the 200 images are "
     "4032×4032, which costs about 15,876 tokens at native resolution; the 4K and 2K budgets are therefore only 26% and 13% of the "
     "native cost. Every input image is resized by us (PIL Lanczos) to width and height that are multiples of 32 before it is passed to "
     "the official processor, so the processor does not resize it again. The token count is measured from the processor's "
     "image_grid_thw as N = t·h·w/4 (the number of <|image_pad|> tokens). An image sent twice within an instance is counted twice, "
     "and no KV cache is reused across calls. The dataset's category and answer fields are used only for scoring and analysis.")
doc.add_heading("2.2 Inference pipeline", level=2)
fig("report/figs/pipeline.png", 17.0, "Figure 1. Final pipeline. Only the blue boxes consume visual tokens.")
bullet(" The whole image is resized to 1,024 tokens (4032² → 1024² px) and the model is prompted to output the option letter first "
       "and then the bounding box of the evidence region. The A–D distribution p1 is the softmax over the four letter logits at the "
       "first decoding step. If max p1 ≥ τ, a stopping criterion ends generation right after that token. Otherwise the bbox_2d "
       "(0–1000 relative coordinates) produced later in the same generation is used, so locating the region costs no extra visual tokens.",
       "Call 1 – Joint probe.")
bullet(" Only for uncertain instances, Qwen sees the question and options without any image and decides whether the question needs "
       "(A) a fine detail of one object/region or (B) relations among multiple objects, counting, or scene layout. The instance is "
       "routed to A if P(A) ≥ 0.3 at the first token.", "Call 2 – Type classification (0 visual tokens).")
bullet(" Type A: the probe box is expanded 2× (minimum side = 25% of the short image side, ≈1,008 px), cropped from the original "
       "image, and fed at ≤1,024 tokens (Success 1 additionally allows up to 2× up-scaling), with the prefix 'This image is a "
       "zoomed-in crop of a larger image.' Type B: Success 1 re-observes the whole image at 4,096 tokens; Success 2 keeps the probe "
       "answer and makes no further call.", "Call 3 – Type-specific escalation.")
bullet(" Final answer = argmax(α·log p1 + log p_esc), with α = 0.25 for type A and α = 0 for type B.", "Fusion.")
para("Joint-probe prompt: ", size=9.5).add_run(
    "\"{question}\\nA. …\\nB. …\\nC. …\\nD. …\\nAnswer with the option's letter from the given choices directly."
    + JOINT_SUFFIX.replace('"', '\\"') + "\"").font.size = Pt(8.5)
para("Type prompt: ", size=9.5).add_run(TYPE_PROMPT.replace("\n", " / ")).font.size = Pt(8.5)
table(["Config", "Probe", "τ", "Type thr.", "Type A (crop)", "Type B", "α (A / B)"],
      [["Success 1", "1,024 tok", "0.995", "0.3", "≤1,024 tok, 2× expand, up-scale ≤2", "global 4,096 tok", "0.25 / 0"],
       ["Success 2", "1,024 tok", "0.999", "0.3", "≤1,024 tok, 2× expand", "keep probe answer", "0.25 / –"]],
      widths=[2.0, 1.8, 1.4, 1.6, 5.0, 2.8, 1.8], size=8.5)
caption("Table 2. Final hyper-parameters (configs/success1_4k.yaml, configs/success2_2k.yaml).")
para("Configuration search. Each stage (probe, several crop variants, global re-observation) was first run once on all 800 instances "
     "without early stopping, logging per-instance A–D probabilities, boxes and token counts. Since later stages do not depend on the "
     "stopping threshold, the accuracy and exact mean token cost of any (τ, threshold, crop variant, α) combination can be computed "
     "offline without re-running the model. The selected configurations were then verified with full end-to-end runs (offline "
     "74.00% / 72.00% → end-to-end 74.00% / 71.88%).")

# ---------------------------------------------------------------- 3
doc.add_heading("3. Experimental setup", level=1)
para("One RTX 3090 (24 GB) per run, Python 3.11, PyTorch 2.5.1 (CUDA 12.4), transformers 4.57.1, bf16, SDPA attention, batch size 1. "
     "Every call uses the original Qwen/Qwen3-VL-2B-Instruct checkpoint and its official AutoProcessor. Decoding is greedy "
     "(do_sample=False, num_beams=1); thinking mode and beam search are not used. Multiple-choice answers are read from the A–D "
     "logits at the first decoding step, so there are no format errors. Latency is the end-to-end wall-clock time per instance, "
     "including image decoding and resizing. Every instance is processed independently; nothing is shared across the four orderings "
     "of the same question.")

# ---------------------------------------------------------------- 4
doc.add_heading("4. Results", level=1)
rows = [run_row(f"Global resize {t}", R[f"global_{t}"]) for t in (1024, 2048, 4096)]
rows += [run_row("Global 512 + 2×2 tiles (384 tok each)", R["tiling_2048"]),
         run_row("Global 1024 + 2×2 tiles (768 tok each)", R["tiling_4096"]),
         run_row("Success 2 (ours)", R["success2_2k"]), run_row("Success 1 (ours)", R["success1_4k"])]
table(["Method", "Mean visual tok", "Acc", "Single", "Cross", "Mean calls"], rows,
      widths=[6.2, 2.6, 1.6, 1.6, 1.6, 1.8], size=8.5, bold_rows=(5, 6))
caption("Table 3. Comparison with baselines (HR-Bench 4K, 800 instances; Single/Cross have 400 instances each).")
para(f"Global down-sampling gains only 3–4 points per doubling of tokens and reaches {R['global_4096']['acc']:.1f}% at 4,096 tokens, "
     f"below the 71% main target. Question-agnostic 2×2 tiling is even worse at the same budget ({R['tiling_2048']['acc']:.1f}% / "
     f"{R['tiling_4096']['acc']:.1f}%): each tile gains little resolution while splitting the scene, which hurts cross-instance questions "
     f"most. Our method is {R['success2_2k']['acc'] - R['global_2048']['acc']:+.1f} points above global-2048 while using only "
     f"{R['success2_2k']['mean_visual']:.0f} tokens, and {R['success1_4k']['acc'] - R['global_4096']['acc']:+.1f} points above global-4096 "
     f"with {R['success1_4k']['mean_visual']:.0f} tokens. For reference, native resolution (≈15,800 tokens) reached "
     f"{S['native_partial']['acc']:.1f}% on a partial run of 54 instances.")
fig("report/figs/acc_tokens.png", 12.0, "Figure 2. Accuracy vs. mean visual tokens. Dashed curves vary only τ of the final policies "
    "(offline); dotted lines mark budgets and main targets.")

doc.add_heading("4.1 Per-call cost and token distribution", level=2)
pc_rows = []
for name, F in [("S1", F1), ("S2", F2)]:
    for tag, lab in [("probe", "Call 1 joint probe"), ("type", "Call 2 type classifier (text-only)"),
                     ("single", "Call 3 crop (type A)"), ("relational", "Call 3 global 4096 (type B)")]:
        if tag in F["per_call"]:
            d = F["per_call"][tag]
            pc_rows.append([name, lab, d["count"], f1(d["mean_visual"]), f1(d["mean_input"] - d["mean_visual"]),
                            f"{d['mean_gen']:.1f}", f"{d['mean_latency']:.3f}"])
table(["Config", "Call", "Count (/800)", "Visual tok", "Text input tok", "Gen. tok", "Latency (s)"], pc_rows,
      widths=[1.2, 5.2, 1.9, 1.8, 2.2, 1.6, 1.9], size=8.5)
caption("Table 4. Mean cost per call type (text input tokens include the vision start/end special tokens).")
para(f"An instance makes {R['success1_4k']['mean_calls']:.2f} (S1) / {R['success2_2k']['mean_calls']:.2f} (S2) Qwen calls and generates "
     f"{F1['gen_tokens_mean']:.1f} / {F2['gen_tokens_mean']:.1f} tokens on average; most generated tokens are the bbox JSON of uncertain "
     f"instances. End-to-end latency is {F1['latency']['mean']:.2f} s mean / {F1['latency']['p95']:.2f} s p95 (S1) and "
     f"{F2['latency']['mean']:.2f} s / {F2['latency']['p95']:.2f} s (S2). Per-instance cost concentrates at 1,024 (early stop), ≈2,048 "
     f"(crop) and ≈5,100 tokens (global re-observation), with maxima of {F1['tokens']['max']} (S1) and {F2['tokens']['max']} (S2) "
     "(Figure 3): tokens saved on easy instances are spent on hard ones.")
fig("report/figs/token_hist.png", 13.0, "Figure 3. Distribution of cumulative visual tokens per instance (black: mean, red dashed: budget).")

# ---------------------------------------------------------------- 5
doc.add_heading("5. Key design decisions and ablation", level=1)
sep, full2 = A2["separate grounding call (+1024 tok)"], A2["Full (Success 2 config)"]
bullet(f" On the 1,024-token probe, the {len(hi)} instances with max p1 ≥ 0.999 are {hi_acc:.1f}% correct, whereas the remaining "
       f"{len(lo)} are only {lo_acc:.1f}% correct, so confidence is a reliable signal for deciding when to look closer.",
       "① Confidence is well calibrated.")
bullet(f" Our first cascade used a separate grounding call after the probe, spending 1,024 tokens only to find the region. Asking for "
       f"the bbox right after the answer letter removes that call: under the same policy it saves {sep['tok'] - full2['tok']:.0f} tokens "
       f"and changes accuracy by {full2['acc'] - sep['acc']:+.2f} points.", "② The joint probe makes localization free.")
bullet(f" Zooming fixed {eff['single'][0]} and broke only {eff['single'][1]} single-instance answers, but on cross-instance questions it "
       f"fixed {eff['cross'][0]} while breaking {eff['cross'][1]}. These questions (relative position, counting, map reasoning) need the "
       "whole scene, which motivated the image-free type router.", "③ Zooming hurts cross-instance questions.")
bullet(f" With the 4K budget, type-B instances are re-observed globally at 4,096 tokens, raising cross accuracy from "
       f"{A1['relational: keep probe']['cross']:.1f}% to {A1['Full (Success 1 config)']['cross']:.1f}%.",
       "④ Type B benefits from a higher-resolution global view.")
abl = [abl_row("Success 2 (full)", full2),
       abl_row("  − early stopping (τ = 1)", A2["w/o early stopping (tau=1)"]),
       abl_row("  − type routing (crop for all)", A2["w/o type routing (zoom all)"]),
       abl_row("  − escalation (joint probe only)", A2["w/o escalation (joint probe only)"]),
       abl_row("  separate grounding call instead of joint", sep),
       abl_row("Success 1 (full)", A1["Full (Success 1 config)"]),
       abl_row("  type B: keep probe (no re-observation)", A1["relational: keep probe"]),
       abl_row("  type A: no up-scaling", A1["single: no 2x up-scaling"]),
       abl_row("  type A: crop at 2,048 tok", A1["single: crop 2048 tok"])]
table(["Configuration (one factor changed)", "Mean tok", "Acc", "Single", "Cross"], abl,
      widths=[7.2, 2.2, 1.8, 1.8, 1.8], size=8.5, bold_rows=(0, 5))
caption("Table 5. Controlled ablations computed from the same force-all logs (identical probe outputs and crops), changing one factor at a time.")
para(f"Escalation is the largest contributor (probe only: {A2['w/o escalation (joint probe only)']['acc']:.2f}%). Early stopping costs "
     f"only {A2['w/o early stopping (tau=1)']['acc'] - full2['acc']:.2f} points while saving "
     f"{A2['w/o early stopping (tau=1)']['tok'] - full2['tok']:.0f} tokens per instance. Type routing keeps 2K accuracy while saving "
     f"{A2['w/o type routing (zoom all)']['tok'] - full2['tok']:.0f} tokens. Raising the crop to 2,048 tokens does not help, so a "
     "1,024-token crop is sufficient.")

# ---------------------------------------------------------------- 6
doc.add_heading("6. Success and failure analysis", level=1)
tx_keys = [("stopped_correct", "Early stop – correct"), ("premature_stop(wrong & confident)", "Early stop – wrong (premature stopping)"),
           ("escalated: fixed (probe wrong -> final right)", "Escalated – fixed (probe wrong → final correct)"),
           ("escalated: kept correct", "Escalated – stayed correct"),
           ("escalated: harmful (probe right -> final wrong)", "Unnecessary escalation (probe correct → final wrong)"),
           ("escalated: still wrong", "Escalated – still wrong")]
table(["Outcome", "Success 1", "Success 2"], [[lab, F1["taxonomy"].get(k, 0), F2["taxonomy"].get(k, 0)] for k, lab in tx_keys],
      widths=[9.0, 2.5, 2.5], size=8.5)
caption("Table 6. Instance-level outcome breakdown of the final runs.")
rv1 = F1["route_vs_category"]
para(f"Premature stopping is rare ({F1['taxonomy'].get('premature_stop(wrong & confident)', 0)} in S1, "
     f"{F2['taxonomy'].get('premature_stop(wrong & confident)', 0)} in S2); a typical case is a book-title question (#216) answered "
     f"wrongly with probability 1.0. Unnecessary escalation affects {F2['taxonomy'].get('escalated: harmful (probe right -> final wrong)', 0)} "
     f"instances in S2 but {F1['taxonomy'].get('escalated: harmful (probe right -> final wrong)', 0)} in S1, because S1 escalates more "
     "(lower τ) and does not fuse the probe for type B. Measured against the dataset labels (analysis only), the router sent "
     f"{rv1.get('cross|single-route', 0)} cross questions to type A and {rv1.get('single|relational-route', 0)} single questions to type B in S1. "
     "Many of the former, such as identifying the president inside a red box on Mount Rushmore, really do need only one region and "
     "were answered correctly by the crop. The remaining errors are mostly cross-instance: map questions asking which numbered "
     f"locations belong to the same country (48 instances) were answered correctly only {QX['s1_map']} (S1) / {QX['s2_map']} (S2) times, "
     f"below chance, and counting questions (40 instances) {QX['s1_cnt']} / {QX['s2_cnt']} times.")
fig("report/figs/qual.png", 16.0, "Figure 4. Representative cases (Success 2). Top: global view with probe box (red) and crop region "
    "(yellow, dashed). Bottom: the crop given to Call 3.")
para("In #252 (license plate) and #344 (speed-limit sign), text that is unreadable in the 1,024-token global view becomes readable "
     "in the native-resolution crop, turning a wrong probe answer into a correct one. #88 combines unreadable detail with confusion: "
     "the box points near a bench where a person and the bicycle are, but the bicycle is only a few dozen pixels wide even at native "
     "resolution, and the model answers with the color of the person's red clothing. #148 is a recognition failure with the right "
     "region: a white object that appears to be the camera is on the pillar, yet the 2B model chooses 'there is no Surveillance "
     "camera'. Because the pipeline relies on a single box, a wrong first localization cannot be recovered.")

# ---------------------------------------------------------------- 7
doc.add_heading("7. Limitations and conclusion", level=1)
bullet(" τ, the type threshold, crop settings and α were selected on the same 800 instances without a separate validation set, and "
       "the type prompt was checked against the category labels. Accuracy varies by about one point for τ ∈ [0.99, 1], so the choice is "
       "not very sensitive, but the reported numbers may be slightly optimistic.", "Tuning on the test set.")
bullet(f" Only <|image_pad|> tokens are counted. Adding the two vision start/end tokens per image (on average {imgs1:.2f} images in S1 "
       f"and {imgs2:.2f} in S2) adds only {2 * imgs1:.1f} / {2 * imgs2:.1f} tokens, which keeps both runs within budget.", "Token accounting.")
para("In summary, obtaining the model's own answer confidence and evidence box from one low-resolution call, and choosing how to look "
     "further with an image-free type decision, reaches "
     f"{R['success1_4k']['acc']:.2f}% (Success 1) and {R['success2_2k']['acc']:.2f}% (Success 2) with fewer tokens than plain global "
     "down-sampling. The main remaining weakness is cross-instance reasoning, especially map questions that require reasoning over the "
     "whole scene.")

out = "report/project1_report.docx"
doc.save(out)
print("saved", out)
